import ast
import glob
import os
import random

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# Ustawienia eksperymentu
NUMBER_OF_RUNS = 100
ITERATIONS = 10000
BITS_PER_DIMENSION = 16
LOWER_BOUND = -10.0
UPPER_BOUND = 10.0

OUTPUT_EXCEL = "Wyniki_LocalSearch.xlsx"


# --------------------------------------------------
# Gray Code i funkcja celu
# --------------------------------------------------

def gray_bits_to_int(bits):
    value = 0
    binary_bit = 0

    for bit in bits:
        binary_bit ^= bit
        value = (value << 1) | binary_bit

    return value


def int_to_gray_bits(value):
    gray_value = value ^ (value >> 1)

    return [
        (gray_value >> i) & 1
        for i in range(BITS_PER_DIMENSION - 1, -1, -1)
    ]


def encode_dimension(value):
    max_integer = (1 << BITS_PER_DIMENSION) - 1

    ratio = (
        (value - LOWER_BOUND) / (UPPER_BOUND - LOWER_BOUND)
    )

    integer_value = round(ratio * max_integer)

    return int_to_gray_bits(integer_value)


def encode_solution(values):
    bits = []

    for value in values:
        bits.extend(encode_dimension(value))

    return bits


def objective_value(bits, dimensions):
    max_integer = (1 << BITS_PER_DIMENSION) - 1
    result = 0.0

    for dimension in range(dimensions):
        start = dimension * BITS_PER_DIMENSION
        end = start + BITS_PER_DIMENSION

        integer_value = gray_bits_to_int(bits[start:end])

        x = LOWER_BOUND + (
            integer_value / max_integer
        ) * (UPPER_BOUND - LOWER_BOUND)

        result += x * x

    return result


# --------------------------------------------------
# Generator sąsiedztwa taki jak w Twoim programie
# --------------------------------------------------

def gray_coding_neighbor_random(bits, m=1):
    neighbor = bits.copy()

    for j in range(len(neighbor)):
        if random.random() < m / 16:
            neighbor[j] = 1 - neighbor[j]

    return neighbor


# --------------------------------------------------
# Jedno uruchomienie Local Search
# --------------------------------------------------

def run_local_search(dimensions):
    x = encode_solution([10.0] * dimensions)
    values = []

    no_improvement_calls = 0
    no_improvement_limit = ITERATIONS

    for _ in range(ITERATIONS):
        candidate = gray_coding_neighbor_random(x, 1)

        if (
            objective_value(candidate, dimensions)
            < objective_value(x, dimensions)
        ):
            x = candidate
            no_improvement_calls = 0
        else:
            no_improvement_calls += 1

        values.append(objective_value(x, dimensions))

        if no_improvement_calls >= no_improvement_limit:
            break

    # Zachowaj dokładnie ITERATIONS wartości.
    # Po wcześniejszym zatrzymaniu stan pozostaje bez zmian.
    if len(values) < ITERATIONS:
        values.extend(
            [values[-1]] * (ITERATIONS - len(values))
        )

    return values


# --------------------------------------------------
# Wczytanie dotychczasowych plików dla n=10
# --------------------------------------------------

def load_existing_n10():
    files = sorted(glob.glob("raw_*.txt"))

    if len(files) != NUMBER_OF_RUNS:
        raise ValueError(
            f"Znaleziono {len(files)} plikow raw_*.txt, "
            f"a oczekiwano {NUMBER_OF_RUNS}.\n"
            "Umiesc w tym folderze dokladnie 100 "
            "dotychczasowych plikow dla n=10."
        )

    runs = []

    print("\nWczytywanie dotychczasowych wynikow n=10...")

    for file_index, filename in enumerate(files, start=1):
        values = []

        with open(filename, "r", encoding="utf-8") as file:
            for line in file:
                line = line.strip()

                if not line:
                    continue

                if line.startswith("["):
                    # Stary format: wektor bitow.
                    bits = ast.literal_eval(line)
                    value = objective_value(bits, 10)
                else:
                    # Obsluga pliku zawierajacego same liczby.
                    value = float(line)

                values.append(value)

        if not values:
            raise ValueError(
                f"Plik {filename} jest pusty."
            )

        if len(values) != ITERATIONS:
            raise ValueError(
                f"Plik {filename} zawiera {len(values)} "
                f"iteracji zamiast {ITERATIONS}."
            )

        runs.append(values)

        if file_index % 10 == 0:
            print(f"Wczytano {file_index}/100 plikow.")

    print("Wczytano wszystkie wyniki n=10.")

    return runs


# --------------------------------------------------
# Obliczanie sredniej
# --------------------------------------------------

def calculate_average(runs):
    number_of_steps = len(runs[0])

    return [
        sum(run[step] for run in runs) / len(runs)
        for step in range(number_of_steps)
    ]


# --------------------------------------------------
# Zapis danych i wykresu do arkusza Excel
# --------------------------------------------------

def add_excel_sheet(workbook, dimensions, runs):
    sheet = workbook.create_sheet(f"n={dimensions}")

    average = calculate_average(runs)

    headers = (
        ["Iteracja"]
        + [
            f"Uruchomienie_{i + 1}"
            for i in range(len(runs))
        ]
        + ["Srednia"]
    )

    sheet.append(headers)

    for step in range(ITERATIONS):
        row = (
            [step + 1]
            + [run[step] for run in runs]
            + [average[step]]
        )

        sheet.append(row)

    # Formatowanie naglowkow.
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4",
        )
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions

    sheet.column_dimensions["A"].width = 12

    for col in range(2, len(runs) + 3):
        sheet.column_dimensions[
            get_column_letter(col)
        ].width = 18

    # Wykres korzysta z kolumny Srednia.
    chart = LineChart()
    chart.title = (
        f"Sredni przebieg Local Search (n={dimensions})"
    )
    chart.style = 13
    chart.y_axis.title = "Srednia wartosc funkcji celu (skala logarytmiczna)"
    chart.y_axis.scaling.logBase = 10
    chart.x_axis.title = "Iteracja"
    chart.height = 9
    chart.width = 20

    average_column = len(runs) + 2

    data = Reference(
        sheet,
        min_col=average_column,
        min_row=1,
        max_row=ITERATIONS + 1,
    )

    categories = Reference(
        sheet,
        min_col=1,
        min_row=2,
        max_row=ITERATIONS + 1,
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    sheet.add_chart(chart, "D2")

    # Dodatkowy wykres PNG.
    plt.figure(figsize=(10, 5))
    plt.semilogy(
        range(1, ITERATIONS + 1),
        average,
    )
    plt.xlabel("Iteracja")
    plt.ylabel("Srednia wartosc funkcji celu (skala logarytmiczna)")
    plt.title(
        f"Average Local Search Progress (n={dimensions})"
    )
    plt.grid(True, which="both", alpha=0.3)
    plt.tight_layout()

    plt.savefig(
        f"wykres_n{dimensions}.png",
        dpi=150,
    )

    plt.close()

    print(f"Utworzono arkusz n={dimensions}.")
    print(
        f"  Srednia na poczatku: {average[0]:.6f}"
    )
    print(
        f"  Srednia na koncu:   {average[-1]:.6f}"
    )


# --------------------------------------------------
# Glowny program
# --------------------------------------------------

def main():
    # Kontrola istnienia starych danych.
    old_files = glob.glob("raw_*.txt")

    if len(old_files) != NUMBER_OF_RUNS:
        print(
            f"Blad: znaleziono {len(old_files)} plikow "
            "raw_*.txt zamiast 100."
        )
        print(
            "Uruchom program z folderu zawierajacego "
            "Twoje dotychczasowe pliki raw dla n=10."
        )
        return

    print("======================================")
    print("   EKSPERYMENT LOCAL SEARCH")
    print("======================================")
    print("Liczba uruchomien na wymiar: 100")
    print("Liczba iteracji: 10000")
    print("Wymiary: 2, 5, 10")
    print()
    print("Istniejace wyniki n=10 zostana wykorzystane.")
    print("Nowe eksperymenty zostana wykonane dla n=2 i n=5.")

    # Wczytaj stare dane przed generowaniem nowych.
    runs_n10 = load_existing_n10()

    # Wykonaj brakujace eksperymenty.
    all_results = {}

    for dimensions in (2, 5):
        runs = []

        print(
            f"\nGenerowanie wynikow dla n={dimensions}..."
        )

        for run_number in range(1, NUMBER_OF_RUNS + 1):
            result = run_local_search(dimensions)
            runs.append(result)

            if run_number % 10 == 0:
                print(
                    f"n={dimensions}: "
                    f"{run_number}/100 uruchomien"
                )

        all_results[dimensions] = runs

    all_results[10] = runs_n10

    # Zbuduj plik Excel.
    workbook = Workbook()
    workbook.remove(workbook.active)

    for dimensions in (2, 5, 10):
        add_excel_sheet(
            workbook,
            dimensions,
            all_results[dimensions],
        )

    workbook.save(OUTPUT_EXCEL)

    print("\n======================================")
    print("GOTOWE!")
    print(f"Plik Excel: {OUTPUT_EXCEL}")
    print("Wykresy: wykres_n2.png, wykres_n5.png, wykres_n10.png")
    print("======================================")


if __name__ == "__main__":
    main()
