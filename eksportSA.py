# Eksport wynikow Simulated Annealing do Excela.
# Umiesc ten plik obok SAmain.py i uruchom go.

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from SAmain import experiment, function_1, function_2

NUMBER_OF_RUNS = 100
ITERATIONS = 10000
OUTPUT_EXCEL = "Wyniki_SA.xlsx"


def add_excel_sheet(workbook, name, representation, runs, average):
    sheet_name = name + "_" + representation
    sheet = workbook.create_sheet(sheet_name)

    headers = ["Wywolanie"]
    for i in range(len(runs)):
        headers.append("Uruchomienie_" + str(i + 1))
    headers.append("Srednia")
    sheet.append(headers)

    for i in range(ITERATIONS):
        row = [i + 1]
        for run in runs:
            row.append(run[i])
        row.append(average[i])
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 14
    for i in range(2, NUMBER_OF_RUNS + 3):
        sheet.column_dimensions[get_column_letter(i)].width = 18

    chart = LineChart()
    chart.title = "Sredni przebieg SA - " + sheet_name
    chart.y_axis.title = "Srednia najlepsza wartosc funkcji"
    chart.x_axis.title = "Numer wywolania funkcji oceny"
    chart.height = 9
    chart.width = 20

    average_column = len(runs) + 2
    data = Reference(sheet, min_col=average_column, min_row=1, max_row=ITERATIONS + 1)
    categories = Reference(sheet, min_col=1, min_row=2, max_row=ITERATIONS + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, "CZ2")

    plt.figure(figsize=(10, 5))
    plt.plot(range(1, ITERATIONS + 1), average)
    plt.xlabel("Numer wywolania funkcji oceny")
    plt.ylabel("Srednia najlepsza wartosc funkcji")
    plt.title("Simulated Annealing - " + sheet_name)
    plt.grid(True)
    plt.tight_layout()
    plt.savefig("wykres_SA_" + sheet_name + ".png", dpi=150)
    plt.close()
    print("Utworzono arkusz:", sheet_name)


def main():
    tests = [
        ("F1", function_1, -3, 3, 0.3),
        ("F2", function_2, -32.768, 32.768, 2.0)
    ]

    workbook = Workbook()
    workbook.remove(workbook.active)

    for name, function, lower, upper, step in tests:
        for representation in ["real", "gray"]:
            print("\nEksperyment:", name, representation)
            runs, average, best_values, best_points = experiment(
                function, lower, upper, representation,
                repetitions=NUMBER_OF_RUNS, counter=ITERATIONS, step=step
            )
            add_excel_sheet(workbook, name, representation, runs, average)
            print("Najlepsza wartosc:", min(best_values))
            print("Srednia koncowa:", average[-1])

    workbook.save(OUTPUT_EXCEL)
    print("\nGOTOWE! Zapisano:", OUTPUT_EXCEL)


if __name__ == "__main__":
    main()
